import io

from docx import Document as WordDocument

from src.grabber.base import GrabbedDocument
from src.parsers.base import (
    BaseParser,
    ParsedDocument,
)


class DocxParser(BaseParser):

    def parse(
        self,
        document: GrabbedDocument,
    ) -> ParsedDocument:

        stream = io.BytesIO(
            document.content
        )

        word_document = WordDocument(
            stream
        )

        parts = []

        paragraph_count = 0
        table_count = 0

        for paragraph in word_document.paragraphs:
            text = paragraph.text.strip()

            if not text:
                continue

            parts.append(text)

            paragraph_count += 1

        for table in word_document.tables:
            table_count += 1

            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                values = [
                    value
                    for value in values
                    if value
                ]

                if values:
                    parts.append(
                        " | ".join(values)
                    )

        text = "\n\n".join(
            parts
        ).strip()

        if not text:
            raise ValueError(
                f"No text could be extracted from "
                f"'{document.name}'."
            )

        return ParsedDocument(
            title=document.name,
            text=text,
            document_type="docx",
            mime_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            external_id=document.external_id,
            metadata={
                "parser": "python-docx",
                "paragraph_count": paragraph_count,
                "table_count": table_count,
            },
        )


docx_parser = DocxParser()