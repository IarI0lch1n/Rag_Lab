import io

from pptx import Presentation

from src.grabber.base import GrabbedDocument
from src.parsers.base import (
    BaseParser,
    ParsedDocument,
)


class PptxParser(BaseParser):
    def parse(
        self,
        document: GrabbedDocument,
    ) -> ParsedDocument:

        presentation = Presentation(
            io.BytesIO(document.content)
        )

        slides_text = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            slide_parts = []

            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()

                    if text:
                        slide_parts.append(text)

                if shape.has_table:
                    for row in shape.table.rows:
                        values = [
                            cell.text.strip()
                            for cell in row.cells
                            if cell.text.strip()
                        ]

                        if values:
                            slide_parts.append(
                                " | ".join(values)
                            )

            if slide_parts:
                slides_text.append(
                    f"[Slide {slide_number}]\n"
                    + "\n".join(slide_parts)
                )

        text = "\n\n".join(
            slides_text
        ).strip()

        if not text:
            raise ValueError(
                f"No text could be extracted from "
                f"'{document.name}'."
            )

        return ParsedDocument(
            title=document.name,
            text=text,
            document_type="pptx",
            mime_type=(
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation"
            ),
            external_id=document.external_id,
            metadata={
                "parser": "python-pptx",
                "slide_count": len(
                    presentation.slides
                ),
            },
        )


pptx_parser = PptxParser()